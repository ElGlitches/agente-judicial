import os
from docx import Document
import re
from docx.oxml.ns import qn

DOCS_PATH = "./docs/"

def get_header_docx(selected_doc, data_dict):
    """
    Genera el objeto Document MODIFICANDO el original para mantener tablas y formato.
    Elimina el cuerpo antiguo y deja solo el encabezado (Presuma + Suma + Intro).
    """
    if selected_doc == "Selección Inteligente (Buscar por tema)" or not selected_doc.endswith(".docx"):
        return Document(), ""
        
    try:
        source_path = os.path.join(DOCS_PATH, selected_doc)
        doc = Document(source_path) # Usamos el original para mantener TABLAS
        preview_text = ""
        
        # Datos para reemplazo
        replacements = {
            "MATERIA": data_dict["materia"],
            "MENOR": data_dict["nombre_menor"],
            "DEMANDANTE": data_dict["nombre_dte"],
            "DEMANDADO": data_dict["nombre_ddo"],
            "ABOGADO": data_dict["abogado"],
            "PATROCINANTE": data_dict["abogado"],
            "APODERADO": data_dict["abogado"],
            "R.U.T": [data_dict["rut_dte"], data_dict["abogado"], data_dict["rut_ddo"]], 
            "DOMICILIO": [data_dict["domicilio"]] 
        }
        ruts = list(replacements["R.U.T"])
        ruts2 = list(replacements["R.U.T"]) # Copia para parrafos


        # --- FUNCIÓN HELPER DE REEMPLAZO ---
        def apply_replacements(text, ruts_list):
            text_upper = text.upper()
            found_change = False
            new_text = text
            
            # Reemplazo Genérico
            for key, val in replacements.items():
                if key in ["R.U.T", "DOMICILIO"]: continue
                if key in text_upper and val:
                    pattern = re.compile(rf"({key}.*?:)(.*)", re.IGNORECASE)
                    if pattern.search(new_text):
                        new_text = pattern.sub(rf"\1 \t {val}", new_text)
                        found_change = True
            
            # Reemplazo RUT (Consumiendo lista)
            if "R.U.T" in text_upper and ruts_list:
                pattern = re.compile(r"(\bR\.?\s?U\.?\s?T\.?\s?:?)(.*)", re.IGNORECASE)
                if pattern.search(new_text):
                    new_text = pattern.sub(rf"\1 \t {ruts_list[0]}", new_text)
                    ruts_list.pop(0)
                    found_change = True
            
            return new_text, found_change
            
        # 1. ESTRATEGIA: RECONSTRUCCIÓN TOTAL para Alimentos
        if "ALIMENTOS" in data_dict["materia"].upper():
            anchor_par = None
            paragraphs_to_delete = []
            
            for p in doc.paragraphs:
                if "LOS HECHOS" in p.text.upper():
                    anchor_par = p
                    break
                paragraphs_to_delete.append(p)
            
            if not anchor_par and doc.paragraphs:
                 anchor_par = doc.paragraphs[0] # Fallback
            
            # Borramos párrafos previos
            for p in paragraphs_to_delete:
                p._element.getparent().remove(p._element)
                
            # Borrar tablas previas
            if doc.tables:
                tbl = doc.tables[0]
                tbl._element.getparent().remove(tbl._element)

            # C. CONSTRUIR NUEVO HEADER (INVERSO)
            suma_text = "EN LO PRINCIPAL: DEMANDA DE ALIMENTOS; PRIMER OTROSÍ: ALIMENTOS PROVISORIOS; SEGUNDO OTROSÍ: ACOMPAÑA DOCUMENTOS; TERCER OTROSÍ: SOLICITA OFICIOS; CUARTO OTROSÍ: PATROCINIO Y PODER."
            trib_text = f"S. J. L. DE FAMILIA DE SAN MIGUEL"
            intro_text = (
                     f"{data_dict['nombre_dte'].upper()}, chilena, cédula nacional de identidad número {data_dict['rut_dte']}, "
                     f"con domicilio (según otrosí), en representación legal de su hijo menor "
                     f"{data_dict['nombre_menor'].upper()}, chileno, del mismo domicilio, a S.S. respetuosamente digo:"
                 )
            
            # Lógica de Género para Bridge Text
            term_rel = data_dict.get("rel_ddo", "Padre").lower() # padre / madre
            term_child = "del niño" if data_dict.get("gen_menor", "Niño") == "Niño" else "de la niña"
            term_child_indiv = "individualizado" if data_dict.get("gen_menor", "Niño") == "Niño" else "individualizada"

            bridge_text = (
                f"Que vengo en interponer la presente demanda de {data_dict['materia']} en contra de {data_dict['nombre_ddo'].upper()}, "
                f"cédula nacional de identidad Nº {data_dict['rut_ddo']}, domiciliado en {data_dict['domicilio'].upper()}, "
                f"en su calidad de {term_rel} {term_child} anteriormente {term_child_indiv}, "
                "todo ello en virtud de los fundamentos de hecho y de derecho que paso a exponer:"
            )
            
            p_bridge = anchor_par.insert_paragraph_before(bridge_text)
            p_bridge.alignment = 3 
            
            p_intro = p_bridge.insert_paragraph_before(intro_text)
            p_intro.alignment = 3 
            
            p_trib = p_intro.insert_paragraph_before(trib_text.upper())
            p_trib.insert_paragraph_before("") 
            
            p_suma = p_trib.insert_paragraph_before(suma_text)
            p_suma.insert_paragraph_before("")

            # Insertar Tabla
            table = doc.add_table(rows=0, cols=2)
            table.style = 'Table Grid'
            p_suma._element.addprevious(table._element)
            
            presuma_data = [
                ("PROCEDIMIENTO", "ORDINARIO"),
                ("MATERIA", data_dict["materia"]),
                ("MENOR", data_dict["nombre_menor"]),
                ("DEMANDANTE", data_dict["nombre_dte"].upper()),
                ("R.U.T", data_dict["rut_dte"]),
                ("DOMICILIO", "GENERAL CELEDÓN Nº 10290, DEPTO 202"), 
                ("", ""), 
                ("ABOGADO PATROCINANTE", data_dict["abogado"].upper()),
                ("R.U.T.", data_dict["rut_abogado_val"]),
                ("DOMICILIO", data_dict["domicilio_abogado"].upper()),
                ("APODERADO", data_dict["abogado"].upper()),
                ("R.U.T.", data_dict["rut_abogado_val"]),
                ("", ""), 
                ("DEMANDADO", data_dict["nombre_ddo"].upper()),
                ("R.U.T.", data_dict["rut_ddo"]),
                ("DOMICILIO", data_dict["domicilio"].upper())
            ]
            
            for label, val in presuma_data:
                row_cells = table.add_row().cells
                row_cells[0].text = label
                row_cells[1].text = str(val)

            preview_text += "[HEADER ALIMENTOS RECONSTRUIDO TOTALMENTE]\n"
            
        else:
            # --- LÓGICA GENERICA (Para cualquier otro escrito) ---
            # 1. Reemplazo en Tablas (Presuma existente)
            if doc.tables:
                for table in doc.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            new_text, changed = apply_replacements(cell.text, ruts)
                            if changed:
                                cell.text = new_text

            # 2. Reemplazo en Párrafos (Intro/Cuerpo)
            # Solo iteramos hasta encontrar "EN LO PRINCIPAL" o similar para no tocar el cuerpo que se borrará?
            # En realidad, el cuerpo se borra despues. Pero el encabezado (Intro) son parrafos normales.
            for p in doc.paragraphs:
                # Evitar tocar "EN LO PRINCIPAL" hacia abajo si se va a borrar?
                # La logica de borrado esta abajo.
                # Si reemplazamos en todo el doc, no pasa nada, luego se borra lo de abajo.
                new_text, changed = apply_replacements(p.text, ruts2)
                if changed:
                    p.text = new_text 

        # 3. TRUNCAR DOCUMENTO
        body_element = doc._body._body
        elements_to_remove = []
        found_cutoff = False
        
        for child in body_element.iterchildren():
            if found_cutoff:
                elements_to_remove.append(child)
                continue
            
            try:
                xml_text = child.xpath('.//w:t')
                full_text = "".join([t.text for t in xml_text if t.text]).upper()
                if "LOS HECHOS" in full_text:
                    found_cutoff = True
                    elements_to_remove.append(child)
            except:
                pass
        
        for child in elements_to_remove:
            body_element.remove(child)

        # Separador visual eliminado para producción
        # doc.add_paragraph("\n" + "_"*40 + "\n(Cuerpo generado por IA comienza aquí)\n")
        
        return doc, preview_text
        
    except Exception as e:
        return Document(), f"Error procesando docx: {e}"
